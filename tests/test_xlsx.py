"""엑셀 파일을 직접 읽고 쓰는 부분. openpyxl 없이 도는가.

사내 pypi 미러에 openpyxl 이 없어서 직접 만들었다. 직접 만든 것이니
'엑셀이 진짜로 읽고 쓰는 꼴' 과 맞는지를 남이 확인해 줘야 하는데, 그
역할을 openpyxl 에 맡긴다 -- 여기 개발 환경에는 있으니 그걸 자로 쓴다.

  - openpyxl 이 쓴 파일을 우리가 읽어서 값이 같은가
  - 우리가 쓴 파일을 openpyxl 이 읽어서 값이 같은가

거기에 openpyxl 이 만들어 주지 않는 것(엑셀 본프로그램이 쓰는
sharedStrings, 중간이 빈 칸, 수식 칸)은 zip 을 손으로 만들어 확인한다.
"""
import io
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.input_manage import input_manage as im

openpyxl = pytest.importorskip("openpyxl")


def by_openpyxl(sheets: dict[str, list[list]]) -> bytes:
    """openpyxl 로 .xlsx 를 만든다 (우리 읽기를 확인할 자)."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def read_by_openpyxl(data: bytes) -> dict[str, list[list]]:
    book = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    return {name: [list(r) for r in book[name].iter_rows(values_only=True)]
            for name in book.sheetnames}


# --------------------------------------------------- openpyxl 이 쓴 것 읽기

def test_reads_what_openpyxl_wrote():
    rows = [["품번", "수량", "비고"],
            ["A-100", 12, "정상"],
            ["B-200", 3.5, "재측정"]]
    assert im.xlsx_read(by_openpyxl({"기준": rows})) == {"기준": rows}


def test_keeps_sheet_order_and_names():
    data = by_openpyxl({"가": [["1"]], "나": [["2"]], "다": [["3"]]})
    assert list(im.xlsx_read(data)) == ["가", "나", "다"]


def test_korean_and_xml_special_characters_survive():
    rows = [["이름", "식"], ["측정 <값>", 'a & b "c"'], ["한글", "탭\t끝"]]
    assert im.xlsx_read(by_openpyxl({"s": rows}))["s"] == rows


def test_empty_cells_come_back_as_none():
    rows = [["a", "b", "c"], ["x", None, "z"]]
    assert im.xlsx_read(by_openpyxl({"s": rows}))["s"] == rows


def test_date_cell_is_a_date_not_a_number():
    data = by_openpyxl({"s": [["언제"], [datetime(2026, 9, 20)]]})
    assert im.xlsx_read(data)["s"][1][0] == date(2026, 9, 20)


def test_datetime_with_a_time_keeps_the_time():
    data = by_openpyxl({"s": [["언제"], [datetime(2026, 9, 20, 13, 30)]]})
    assert im.xlsx_read(data)["s"][1][0] == datetime(2026, 9, 20, 13, 30)


def test_booleans_stay_booleans():
    data = by_openpyxl({"s": [["쓰나"], [True], [False]]})
    got = im.xlsx_read(data)["s"]
    assert got[1][0] is True and got[2][0] is False


def test_whole_number_does_not_become_a_float():
    """1 이 1.0 으로 읽히면 화면에도 '1.0' 으로 찍히고 그대로 저장된다."""
    data = by_openpyxl({"s": [["n"], [1], [1000000]]})
    assert im.xlsx_read(data)["s"][1:] == [[1], [1000000]]


# ------------------------------------------- 엑셀 본프로그램이 쓰는 꼴 읽기

def hand_made(sheet_xml: str, shared: list[str] | None = None,
              styles: str | None = None) -> bytes:
    """openpyxl 이 만들어 주지 않는 꼴을 손으로 만든다."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
            ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets></workbook>')
        zf.writestr("xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
            'relationships"><Relationship Id="rId1" Type="x" Target="worksheets/sheet1.xml"/>'
            '</Relationships>')
        zf.writestr("xl/worksheets/sheet1.xml",
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f"<sheetData>{sheet_xml}</sheetData></worksheet>")
        if shared is not None:
            zf.writestr("xl/sharedStrings.xml",
                '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                + "".join(f"<si><t>{t}</t></si>" for t in shared) + "</sst>")
        if styles is not None:
            zf.writestr("xl/styles.xml", styles)
    return buf.getvalue()


def test_reads_shared_strings():
    """엑셀 본프로그램은 글자를 따로 모아 두고 칸에는 번호만 적는다."""
    data = hand_made('<row r="1"><c r="A1" t="s"><v>1</v></c>'
                     '<c r="B1" t="s"><v>0</v></c></row>',
                     shared=["둘째", "첫째"])
    assert im.xlsx_read(data)["Sheet1"] == [["첫째", "둘째"]]


def test_shared_string_split_across_runs_is_joined():
    """글자 가운데만 굵게 하면 <t> 가 쪼개진다. 이어 붙여야 한 낱말이다."""
    buf = io.BytesIO(hand_made('<row r="1"><c r="A1" t="s"><v>0</v></c></row>',
                               shared=["x"]))
    with zipfile.ZipFile(buf) as src:
        parts = {n: src.read(n) for n in src.namelist()}
    parts["xl/sharedStrings.xml"] = (
        b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b"<si><r><t>\xea\xb8\xb0\xec\xa4\x80</t></r><r><t>\xec\xa0\x95\xeb\xb3\xb4</t></r>"
        b"</si></sst>")
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zf:
        for name, body in parts.items():
            zf.writestr(name, body)
    assert im.xlsx_read(out.getvalue())["Sheet1"] == [["기준정보"]]


def test_skipped_cells_keep_their_column():
    """중간 칸이 비면 엑셀은 그 칸을 아예 안 적는다. 자리를 채워야 한다."""
    data = hand_made('<row r="1"><c r="A1" t="str"><v>a</v></c>'
                     '<c r="D1" t="str"><v>d</v></c></row>')
    assert im.xlsx_read(data)["Sheet1"] == [["a", None, None, "d"]]


def test_skipped_rows_keep_their_row_number():
    data = hand_made('<row r="1"><c r="A1" t="str"><v>a</v></c></row>'
                     '<row r="3"><c r="A3" t="str"><v>c</v></c></row>')
    assert im.xlsx_read(data)["Sheet1"] == [["a"], [], ["c"]]


def test_formula_cell_reads_the_last_computed_value():
    data = hand_made('<row r="1"><c r="A1" t="str"><f>SUM(B1:C1)</f><v>7</v></c></row>')
    assert im.xlsx_read(data)["Sheet1"] == [["7"]]


def test_custom_date_format_is_recognised():
    """미리 정해진 번호가 아니라 사람이 만든 날짜 서식이어도 날짜다."""
    styles = ('<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
              '2006/main"><numFmts><numFmt numFmtId="164" formatCode="yyyy-mm-dd"/>'
              '</numFmts><cellXfs count="2">'
              '<xf numFmtId="0"/><xf numFmtId="164"/></cellXfs></styleSheet>')
    data = hand_made('<row r="1"><c r="A1" s="1"><v>46285</v></c>'
                     '<c r="B1" s="0"><v>46285</v></c></row>', styles=styles)
    got = im.xlsx_read(data)["Sheet1"][0]
    assert got[0] == date(2026, 9, 20)
    assert got[1] == 46285          # 서식이 없으면 그냥 수다


def test_quoted_text_in_a_format_is_not_mistaken_for_a_date():
    """'0"m"' 은 단위 m 을 붙이는 수 서식이지 달(m)이 아니다."""
    styles = ('<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
              '2006/main"><numFmts><numFmt numFmtId="164" formatCode="0&quot;m&quot;"/>'
              '</numFmts><cellXfs count="1"><xf numFmtId="164"/></cellXfs></styleSheet>')
    data = hand_made('<row r="1"><c r="A1" s="0"><v>5</v></c></row>', styles=styles)
    assert im.xlsx_read(data)["Sheet1"] == [[5]]


def test_garbage_is_refused_clearly():
    with pytest.raises(im.BadWorkbook):
        im.xlsx_read(b"this is not a zip")


def test_zip_without_sheets_is_refused():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
            '2006/main"><sheets/></workbook>')
    with pytest.raises(im.BadWorkbook):
        im.xlsx_read(buf.getvalue())


# ------------------------------------------------ 우리가 쓴 것을 openpyxl 이

def test_openpyxl_reads_what_we_wrote():
    rows = [["품번", "수량"], ["A-100", 12], ["한글", 3.5]]
    assert read_by_openpyxl(im.xlsx_write({"기준": rows})) == {"기준": rows}


def test_written_file_keeps_sheet_order():
    data = im.xlsx_write({"가": [["1"]], "나": [["2"]], "다": [["3"]]})
    assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames == ["가", "나", "다"]


def test_written_special_characters_survive_openpyxl():
    rows = [["<열> & 이름"], ['따옴표 "안"'], ["한글 & 영문"]]
    assert read_by_openpyxl(im.xlsx_write({"s": rows}))["s"] == rows


def test_control_characters_are_dropped_not_written():
    """제어문자가 그대로 들어가면 엑셀이 파일 자체를 못 연다."""
    data = im.xlsx_write({"s": [["앞\x07뒤"]]})
    assert read_by_openpyxl(data)["s"] == [["앞뒤"]]


def test_round_trip_through_our_own_reader():
    rows = [["a", "b", "c"], ["가", 1, 2.5], ["나", None, "끝"]]
    assert im.xlsx_read(im.xlsx_write({"s": rows}))["s"] == rows


def test_wide_sheet_uses_two_letter_columns():
    rows = [[f"c{i}" for i in range(30)]]
    assert im.xlsx_read(im.xlsx_write({"s": rows}))["s"] == rows
    assert im.col_letter(26) == "AA" and im.col_index("AA1") == 26


# ------------------------------------------------------- DataFrame 까지

def frames(sheets):
    return im.read_xlsx(im.to_xlsx(sheets))


def test_dataframe_round_trip():
    df = pd.DataFrame({"품번": ["A-100", "B-200"], "수량": ["12", "3"]}, dtype=object)
    out = frames({"기준": df})["기준"]
    assert list(out.columns) == ["품번", "수량"]
    assert out.values.tolist() == [["A-100", 12], ["B-200", 3]]


def test_numbers_typed_in_the_grid_are_stored_as_numbers():
    """격자는 모든 값을 글자로 올려보낸다. 수로 적은 것은 수로 저장해야
    엑셀에서 정렬과 합계가 되고, 다시 읽어도 '12' 그대로다."""
    df = pd.DataFrame({"n": ["12", "-3", "1.5"]}, dtype=object)
    assert frames({"s": df})["s"]["n"].tolist() == [12, -3, 1.5]


@pytest.mark.parametrize("text", ["0010", "1.50", "1e5", "01", "+3", " 7", "1,000"])
def test_text_that_only_looks_numeric_stays_text(text):
    """앞의 0 은 품번에서 뜻이 있고 '1.50' 은 사람이 적은 자릿수다.
    저장이 사람 친 글자를 조용히 고치면 안 된다."""
    df = pd.DataFrame({"n": [text]}, dtype=object)
    assert frames({"s": df})["s"]["n"].tolist() == [text]


def test_blank_rows_are_dropped_on_the_way_out():
    df = pd.DataFrame({"a": ["x", "", None], "b": ["y", "", ""]}, dtype=object)
    assert len(frames({"s": df})["s"]) == 1


def test_missing_cells_come_back_as_none_not_nan():
    """'nan' 이라는 글자로 저장되는 것을 막는 자리다."""
    df = pd.DataFrame({"a": ["x", "y"], "b": ["1", None]}, dtype=object)
    out = frames({"s": df})["s"]
    assert out["b"].tolist() == [1, None]


def test_duplicate_column_names_get_numbered():
    """겹친 채로 두면 df[이름] 이 Series 가 아니라 DataFrame 이 되고,
    그때부터 값 대신 표가 실려 나간다."""
    data = im.xlsx_write({"s": [["a", "a", "a"], ["1", "2", "3"]]})
    out = im.read_xlsx(data)["s"]
    assert list(out.columns) == ["a", "a_1", "a_2"]


def test_nameless_column_gets_a_placeholder():
    data = im.xlsx_write({"s": [["a", None, "c"], ["1", "2", "3"]]})
    assert list(im.read_xlsx(data)["s"].columns) == ["a", "Unnamed: 1", "c"]


def test_short_rows_are_padded_to_the_header():
    data = im.xlsx_write({"s": [["a", "b", "c"], ["1"], ["1", "2", "3"]]})
    out = im.read_xlsx(data)["s"]
    assert out.shape == (2, 3)
    assert out.iloc[0].tolist() == ["1", None, None]


def test_row_longer_than_the_header_is_not_thrown_away():
    """머리글보다 값이 더 오른쪽까지 있으면 그 값도 살려야 한다."""
    data = im.xlsx_write({"s": [["a"], ["1", "2"]]})
    out = im.read_xlsx(data)["s"]
    assert out.shape == (1, 2)
    assert out.iloc[0].tolist() == ["1", "2"]


def test_empty_sheet_is_an_empty_frame():
    assert im.read_xlsx(im.xlsx_write({"s": []}))["s"].empty


def test_long_sheet_names_do_not_collide():
    """엑셀 시트 이름은 31자까지다. 잘라서 같아지면 한 시트가 사라진다."""
    a, b = "기" * 30 + "가", "기" * 30 + "나"
    out = im.read_xlsx(im.to_xlsx({a: pd.DataFrame({"c": ["1"]}),
                                   b: pd.DataFrame({"c": ["2"]})}))
    assert len(out) == 2
    assert [df["c"].tolist() for df in out.values()] == [[1], [2]]


def test_eight_thousand_rows_round_trip():
    df = pd.DataFrame({"품번": [f"A-{i}" for i in range(8500)],
                       "수량": [str(i) for i in range(8500)]}, dtype=object)
    out = frames({"기준": df})["기준"]
    assert len(out) == 8500
    assert out.iloc[-1].tolist() == ["A-8499", 8499]


# ------------------------------------------------ 끌어 채운 수식 (공유 수식)

SHIFT_CASES = [
    ("VLOOKUP(E2,ET추출여부!$A:$C,3,0)", 3, 0),
    ("VLOOKUP(E2,ET추출여부!$A:$C,3,0)", 10218, 0),
    ("SUM(A2:B2)*$C$1", 2, 1),
    ('IF(A2="A2",B$2,C2)', 1, 0),
    ("'My Sheet'!A2+LOG10(B2)", 1, 0),
    ("SUM(2:2)+A:A", 1, 1),
    ("$A2+A$2+$A$2+A2", 5, 3),
    ("Sheet2!B7*2", 4, 2),
    ("ATAN2(A2,B2)+DAYS360(C2,D2)", 1, 0),
    ('"x"&A2&"y""z"&B2', 2, 0),
]


@pytest.mark.parametrize("formula,dr,dc", SHIFT_CASES)
def test_shifting_a_formula_matches_what_excel_does(formula, dr, dc):
    """엑셀이 채우기 할 때와 같게 옮기는가 -- openpyxl 의 Translator 를 자로."""
    from openpyxl.formula.translate import Translator
    want = Translator("=" + formula, origin="C2").translate_formula(
        f"{im.col_letter(2 + dc)}{2 + dr}")[1:]
    assert im._shift_formula(formula, dr, dc) == want


def test_a_reference_shifted_off_the_sheet_is_refused():
    with pytest.raises(im._BadShift):
        im._shift_formula("A1+B2", -1, 0)


def shared_sheet():
    """엑셀이 '아래로 끌어 채우기' 한 VLOOKUP 을 저장하는 모양 그대로:
    맨 윗칸(C2)에만 수식이 있고 나머지(C3, C4)는 번호(si)만 가리킨다."""
    return hand_made(
        '<row r="1"><c r="A1" t="inlineStr"><is><t>code</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>x</t></is></c>'
        '<c r="C1" t="inlineStr"><is><t>찾은값</t></is></c></row>'
        '<row r="2"><c r="A2" t="inlineStr"><is><t>K1</t></is></c><c r="B2"><v>1</v></c>'
        '<c r="C2" t="str"><f t="shared" ref="C2:C4" si="0">VLOOKUP(A2,Sheet1!$A:$B,2,0)'
        '</f><v>1</v></c></row>'
        '<row r="3"><c r="A3" t="inlineStr"><is><t>K2</t></is></c><c r="B3"><v>2</v></c>'
        '<c r="C3" t="str"><f t="shared" si="0"/><v>2</v></c></row>'
        '<row r="4"><c r="A4" t="inlineStr"><is><t>K3</t></is></c><c r="B4"><v>3</v></c>'
        '<c r="C4" t="str"><f t="shared" si="0"/><v>3</v></c></row>')


def test_a_filled_down_formula_is_read_into_every_cell():
    """맨 윗칸만 읽으면 저장할 때 나머지 칸의 수식이 값으로 굳는다."""
    got = {}
    im.read_xlsx(shared_sheet(), got)
    assert got["Sheet1"] == {
        (0, "찾은값"): "VLOOKUP(A2,Sheet1!$A:$B,2,0)",
        (1, "찾은값"): "VLOOKUP(A3,Sheet1!$A:$B,2,0)",
        (2, "찾은값"): "VLOOKUP(A4,Sheet1!$A:$B,2,0)",
    }


def test_a_filled_down_formula_survives_a_save():
    formulas = {}
    sheets = im.read_xlsx(shared_sheet(), formulas)
    book = openpyxl.load_workbook(io.BytesIO(im.to_xlsx(sheets, formulas)))
    assert [book["Sheet1"][f"C{r}"].value for r in (2, 3, 4)] == [
        "=VLOOKUP(A2,Sheet1!$A:$B,2,0)",
        "=VLOOKUP(A3,Sheet1!$A:$B,2,0)",
        "=VLOOKUP(A4,Sheet1!$A:$B,2,0)"]


def test_an_array_formula_stays_an_array_formula():
    """배열 수식({=...})을 보통 수식으로 적으면 엑셀이 다르게 계산한다."""
    data = hand_made(
        '<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>합</t></is></c></row>'
        '<row r="2"><c r="A2"><v>1</v></c>'
        '<c r="B2"><f t="array" ref="B2:B3">A2:A3*2</f><v>2</v></c></row>'
        '<row r="3"><c r="A3"><v>2</v></c><c r="B3"><v>4</v></c></row>')
    formulas = {}
    sheets = im.read_xlsx(data, formulas)
    out = im.to_xlsx(sheets, formulas)
    with zipfile.ZipFile(io.BytesIO(out)) as zf:
        xml = zf.read("xl/worksheets/sheet1.xml").decode()
    assert '<f t="array" ref="B2:B3">A2:A3*2</f>' in xml, xml
    assert "<f>" not in xml.split('ref="B2:B3"')[1], "나머지 칸에 수식이 생겼다"


@pytest.mark.parametrize("code,is_date", [
    ("yyyy-mm-dd", True),
    ('[$-412]yyyy"년" m"월" d"일"', True),       # 한국어 엑셀의 날짜
    ("[h]:mm:ss", True),                          # 지난 시간
    ("h:mm AM/PM", True),
    ("#,##0;[Red]-#,##0", False),                 # Red 의 d 는 날짜가 아니다
    ("#,##0_);[Red](#,##0)", False),
    ('_-* #,##0_-;-* #,##0_-;_-* "-"_-;_-@_-', False),   # 회계
    ("[$USD] #,##0.00", False),
    ('0"m"', False),
    ("0.00E+00", False),
    ("@", False),
    ("General", False),
])
def test_what_counts_as_a_date_format(code, is_date):
    assert im._is_date_format(code) is is_date


def test_a_red_negative_format_keeps_numbers_numbers():
    """회계 서식 [Red] 의 d 를 날짜로 보면 그 칸의 수가 전부 날짜로 바뀐다."""
    styles = ('<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
              '2006/main"><numFmts><numFmt numFmtId="164" formatCode="#,##0;[Red]-#,##0"/>'
              '</numFmts><cellXfs count="2">'
              '<xf numFmtId="0"/><xf numFmtId="164"/></cellXfs></styleSheet>')
    data = hand_made('<row r="1"><c r="A1" s="1"><v>1500</v></c></row>', styles=styles)
    assert im.xlsx_read(data)["Sheet1"] == [[1500]]


# ------------------------------------------------ 날짜는 날짜로 남는다

def test_a_date_is_written_as_a_date_not_as_text():
    """글자로 적으면 그 칸이 날짜가 아니게 되어 pandas 로 읽는 쪽의 형이 바뀐다."""
    df = pd.DataFrame({"일자": [date(2026, 9, 20)],
                       "시각": [datetime(2026, 9, 20, 13, 45)]}, dtype=object)
    data = im.to_xlsx({"s": df})
    ws = openpyxl.load_workbook(io.BytesIO(data))["s"]
    assert ws["A2"].is_date and ws["A2"].value == datetime(2026, 9, 20)
    assert ws["B2"].is_date and ws["B2"].value == datetime(2026, 9, 20, 13, 45)
    back = im.read_xlsx(data)["s"]
    assert back.loc[0, "일자"] == date(2026, 9, 20)
    assert back.loc[0, "시각"] == datetime(2026, 9, 20, 13, 45)


def test_dates_that_went_through_the_grid_come_back_as_dates():
    """격자는 모든 값을 글자로 돌려준다. 원래 날짜 칸이면 날짜로 되돌린다."""
    base = pd.DataFrame({"일자": [date(2026, 9, 20), None], "코드": ["a", "b"],
                         "사용": [True, False]}, dtype=object)
    got = im.to_frames({"sheets": [{
        "name": "S", "orig": "S", "cols": ["일자", "코드", "사용"],
        "rows": [["2026-10-01", "2026-10-02", "True"], ["미정", "b", "False"]]}]},
        {"S": base})["S"]
    assert got["일자"].tolist() == [date(2026, 10, 1), "미정"]
    assert got["코드"].tolist() == ["2026-10-02", "b"], "날짜가 없던 칸은 글자로 둔다"
    assert got["사용"].tolist() == [True, False]


@pytest.mark.parametrize("text,want", [
    ("12", True), ("-1.5", True), ("1e5", True), (".5", True),
    ("nan", False), ("inf", False), ("1_000", False), (" 1", False), ("", False),
])
def test_only_what_excel_reads_as_a_number_is_written_as_one(text, want):
    """'nan' 이나 '1_000' 을 수로 적으면 엑셀이 파일을 열며 복구 창을 띄운다."""
    assert im._looks_numeric(text) is want


def test_an_unnamed_header_is_written_back_blank():
    """비어 있던 머리글은 읽을 때 'Unnamed: 1' 로 부르지만 파일에는 빈칸이다."""
    data = hand_made('<row r="1"><c r="A1" t="inlineStr"><is><t>a</t></is></c></row>'
                     '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>2</v></c></row>')
    sheets = im.read_xlsx(data)
    assert list(sheets["Sheet1"].columns) == ["a", "Unnamed: 1"]
    ws = openpyxl.load_workbook(io.BytesIO(im.to_xlsx(sheets)))["Sheet1"]
    assert ws["B1"].value is None and ws["B2"].value == 2


def test_phonetic_guides_are_not_part_of_the_text():
    """일본어 후리가나(<rPh>)는 칸 글자가 아니다."""
    # hand_made 가 <si><t>..</t></si> 로 감싸므로 그 사이에 발음 표기를 끼운다
    data = hand_made('<row r="1"><c r="A1" t="s"><v>0</v></c></row>',
                     shared=['東京</t><rPh sb="0" eb="2"><t>とうきょう</t></rPh><t>'])
    assert im.xlsx_read(data)["Sheet1"] == [["東京"]]
