"""기준 정보가 조용히 사라지거나 덮어써지지 않는가 (S3 는 가짜로 대신한다)."""
import io
import sys
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.input_manage import input_manage as im
from tests import fake_s3


@pytest.fixture(autouse=True)
def fake(monkeypatch):
    fake_s3.reset()
    im._BOOKS.clear()
    monkeypatch.setattr(im, "s3", fake_s3)
    monkeypatch.setattr(im, "FOLDER_PATH", "2GAPU/input")
    return fake_s3


def sheets(**kw):
    return {k: pd.DataFrame(v) for k, v in kw.items()}


def put_book(book, data):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for n, df in data.items():
            df.to_excel(w, sheet_name=n, index=False)
    fake_s3.put_object(f"2GAPU/input/{book}.xlsx", buf.getvalue())


# ---------------------------------------------------------- 목록 / 왕복

def test_only_real_workbooks_are_listed():
    put_book("OCAP기준", sheets(S=[{"a": 1}]))
    put_book("설비기준", sheets(S=[{"a": 1}]))
    fake_s3.put_object("2GAPU/input/_audit.csv", b"x")
    fake_s3.put_object("2GAPU/input/_history/OCAP기준/20260101_000000_a.xlsx", b"x")
    fake_s3.put_object("2GAPU/input/읽어보기.txt", b"x")
    assert im.list_workbooks() == ["OCAP기준", "설비기준"]


def test_a_round_trip_keeps_every_sheet_and_value():
    im.save_workbook("OCAP기준", sheets(
        코드표=[{"code": "C1", "desc": "재작업"}, {"code": "C2", "desc": "폐기"}],
        담당자=[{"item": "item1", "owner": "홍길동"}]), "hong")
    back, _ = im.load_workbook("OCAP기준")
    assert list(back) == ["코드표", "담당자"]
    assert back["코드표"]["desc"].tolist() == ["재작업", "폐기"]


def test_workbooks_do_not_bleed_into_each_other():
    im.save_workbook("A", sheets(S=[{"v": "a"}]), "hong")
    im.save_workbook("B", sheets(S=[{"v": "b"}]), "hong")
    assert im.load_workbook("A")[0]["S"]["v"].tolist() == ["a"]
    assert im.load_workbook("B")[0]["S"]["v"].tolist() == ["b"]


def test_a_missing_workbook_opens_empty_instead_of_raising():
    got, stamp = im.load_workbook("없는파일")
    assert got == {} and stamp == ""


# ------------------------------------------------------ 동시 편집 / 실패

def test_a_second_person_saving_first_is_refused_not_overwritten():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    _mine, stamp = im.load_workbook("A")

    im.save_workbook("A", sheets(S=[{"a": 2}]), "kim")       # 옆 사람이 먼저

    with pytest.raises(im.ConcurrentEdit):
        im.save_workbook("A", sheets(S=[{"a": 3}]), "hong", base_stamp=stamp)
    assert im.load_workbook("A")[0]["S"]["a"].tolist() == [2], \
        "먼저 저장한 값이 덮어써졌습니다"


def test_the_guard_is_per_workbook_not_global():
    """다른 파일을 누가 저장했다고 내 저장이 막히면 안 된다."""
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    _s, stamp_a = im.load_workbook("A")
    im.save_workbook("B", sheets(S=[{"b": 1}]), "kim")
    im.save_workbook("A", sheets(S=[{"a": 2}]), "hong", base_stamp=stamp_a)
    assert im.load_workbook("A")[0]["S"]["a"].tolist() == [2]


def test_saving_with_the_returned_stamp_goes_through():
    stamp = im.save_workbook("A", sheets(S=[{"a": 1}]), "hong").stamp
    stamp = im.save_workbook("A", sheets(S=[{"a": 2}]), "hong",
                             base_stamp=stamp).stamp
    im.save_workbook("A", sheets(S=[{"a": 3}]), "hong", base_stamp=stamp)
    assert im.load_workbook("A")[0]["S"]["a"].tolist() == [3]


def test_opening_the_same_version_twice_reads_it_once(monkeypatch):
    """8MB 짜리를 읽는 데 몇 초다. 판이 같으면 다시 읽을 까닭이 없다."""
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    reads, gets = [], []
    real_read, real_get = im.read_xlsx, fake_s3.get_object
    monkeypatch.setattr(im, "read_xlsx",
                        lambda *a, **k: reads.append(1) or real_read(*a, **k))
    monkeypatch.setattr(fake_s3, "get_object",
                        lambda key: gets.append(key) or real_get(key))
    first = im.fetch_workbook("A")
    again = im.fetch_workbook("A")
    assert len(reads) == 1, "같은 판을 두 번 읽었습니다"
    assert len(gets) == 1, "같은 판인데 파일을 또 내려받았습니다"
    assert again[1]["S"]["a"].tolist() == first[1]["S"]["a"].tolist() == [1]


def test_someone_elses_save_is_read_fresh():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    assert im.fetch_workbook("A")[1]["S"]["a"].tolist() == [1]
    im.save_workbook("A", sheets(S=[{"a": 2}]), "kim")       # 판이 바뀐다
    raw, got, _f, stamp = im.fetch_workbook("A")
    assert got["S"]["a"].tolist() == [2], "지난 판을 그대로 보여줍니다"
    assert stamp == fake_s3.head_etag("2GAPU/input/A.xlsx")
    assert raw == fake_s3.STORE["2GAPU/input/A.xlsx"]


def test_a_file_that_is_gone_opens_empty_and_forgets_the_old_copy():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    im.fetch_workbook("A")
    fake_s3.delete_object("2GAPU/input/A.xlsx")
    assert im.fetch_workbook("A") == (b"", {}, {}, "")


def test_the_gc_is_back_on_after_reading():
    import gc
    assert gc.isenabled()
    im.read_xlsx(im.to_xlsx(sheets(S=[{"a": 1}])))
    assert gc.isenabled(), "읽고 나서 GC 를 다시 안 켰습니다"


def test_the_gc_is_back_on_even_when_reading_fails():
    import gc
    with pytest.raises(im.BadWorkbook):
        im.read_xlsx(b"not a zip")
    assert gc.isenabled()


def test_the_gc_stays_off_if_it_was_off_before():
    """누가 일부러 꺼 두었으면 우리가 켜면 안 된다."""
    import gc
    gc.disable()
    try:
        im.read_xlsx(im.to_xlsx(sheets(S=[{"a": 1}])))
        assert not gc.isenabled()
    finally:
        gc.enable()


def test_what_a_save_hands_back_is_what_landed_in_s3():
    """저장한 뒤 화면은 이걸로 맞춘다 -- S3 에서 도로 안 읽으려고."""
    done = im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    assert done.body == fake_s3.STORE["2GAPU/input/A.xlsx"]
    assert done.stamp == fake_s3.head_etag("2GAPU/input/A.xlsx")
    assert im.read_xlsx(done.body)["S"]["a"].tolist() == [1]


def test_a_failed_save_leaves_the_stored_file_intact():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    before = fake_s3.STORE["2GAPU/input/A.xlsx"]

    def explode(*_a, **_k):
        raise RuntimeError("S3 가 응답하지 않는다")

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(im, "xlsx_write", explode)
        with pytest.raises(Exception):
            im.save_workbook("A", sheets(S=[{"a": 2}]), "hong")
    assert fake_s3.STORE["2GAPU/input/A.xlsx"] == before
    assert im.load_workbook("A")[0]["S"]["a"].tolist() == [1]


# ------------------------------------- 이력은 REV_INFO 시트 한 곳에만

def test_the_only_things_written_are_the_workbook_and_its_copy():
    """감사 기록은 옆 파일이 아니라 엑셀 안의 REV_INFO 시트에 쌓인다.

    옆에 파일을 두면 그 둘이 갈리기 시작한다. 되돌릴 사본은 다른 얘기라
    이력 폴더에 따로 쌓는다.
    """
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    im.save_workbook("A", sheets(S=[{"a": 2}]), "kim")
    rest = [k for k in fake_s3.STORE
            if k != "2GAPU/input/A.xlsx" and "/이력/" not in k]
    assert rest == [], rest


def test_the_change_log_lands_in_the_workbook_itself():
    book = rev_book()
    after = {k: v.copy() for k, v in book.items()}
    after["STEP"].loc[0, "b"] = "바뀜"
    changes = im.workbook_changes(book, after)
    body = im.append_rev_info(after, "2026-09-21", "사유", "나", "",
                              im.changes_text(changes))
    im.save_workbook("A", body, "나")

    back = im.load_workbook("A")[0]["REV_INFO"]
    assert back.iloc[-1]["Remark"] == "사유"
    assert "수정 1행" in back.iloc[-1]["관련"], back.iloc[-1]["관련"]


# ------------------------------------------------------------ 값 다듬기

def test_blank_rows_added_by_the_grid_are_not_saved():
    im.save_workbook("A", sheets(
        S=[{"a": "x", "b": "y"}, {"a": None, "b": None}, {"a": "", "b": "  "}]), "hong")
    assert len(im.load_workbook("A")[0]["S"]) == 1


def test_missing_values_are_saved_blank_not_as_the_word_nan():
    im.save_workbook("A", sheets(S=[{"a": "x", "b": None}]), "hong")
    assert im.load_workbook("A")[0]["S"]["b"].tolist() == [None]


@pytest.mark.parametrize("before,after,want", [
    ([{"a": 1}], [{"a": 1}], 0),
    ([{"a": 1}], [{"a": 2}], 1),
    ([{"a": 1}], [{"a": 1}, {"a": 2}], 1),              # 줄이 늘었다
    ([{"a": 1}, {"a": 2}], [{"a": 1}], 1),              # 줄이 줄었다
    ([{"a": 1}], [{"a": 1, "b": 2}], 1),                # 칸이 늘었다
])
def test_changed_cells_counts_what_a_person_would_count(before, after, want):
    assert im.changed_cells(pd.DataFrame(before), pd.DataFrame(after)) == want


@pytest.mark.parametrize("name", ["아주긴시트이름" * 10, "재고/현황", "a[b]c", "매출:합"])
def test_sheet_names_excel_would_refuse_do_not_break_the_save(name):
    im.save_workbook("A", {name: pd.DataFrame([{"a": 1}])}, "hong")
    got = list(im.load_workbook("A")[0])[0]
    assert len(got) <= 31 and not set(got) & set(':\\/?*[]'), got


def test_adding_an_empty_column_counts_as_a_change():
    """값만 견주면 0 이 나와서 열을 넣고 저장을 누를 수가 없다."""
    before = pd.DataFrame([{"a": "1"}, {"a": "2"}])
    after = pd.DataFrame([{"a": "1", "새칸": ""}, {"a": "2", "새칸": ""}])
    assert im.changed_cells(before, after) == 1


def test_removing_an_empty_column_counts_as_a_change():
    before = pd.DataFrame([{"a": "1", "빈칸": ""}])
    after = pd.DataFrame([{"a": "1"}])
    assert im.changed_cells(before, after) == 1


def test_a_column_added_with_content_is_not_counted_twice():
    """내용이 있으면 그 값이 이미 세어졌으므로 더하지 않는다."""
    before = pd.DataFrame([{"a": "1"}])
    after = pd.DataFrame([{"a": "1", "b": "2"}])
    assert im.changed_cells(before, after) == 1


# -------------------------------------- 저장 전에 보여 줄 것 (무엇이 바뀌나)

def frame(rows, cols=("a", "b")):
    return pd.DataFrame(rows, columns=list(cols), dtype=object)


def kinds(before, after):
    rows, total = im.row_changes(before, after)
    return [(r["kind"], r["row"]) for r in rows], total


def test_an_edited_row_is_called_edited():
    before = frame([["1", "x"], ["2", "y"]])
    after = frame([["1", "x"], ["2", "바뀜"]])
    assert kinds(before, after) == ([("수정", 2)], 1)


def test_a_row_added_at_the_end_is_called_new():
    before = frame([["1", "x"]])
    after = frame([["1", "x"], ["2", "y"]])
    assert kinds(before, after) == ([("신규", 2)], 1)


def test_a_row_inserted_in_the_middle_does_not_mark_the_rest_as_edited():
    """자리만 맞춰 비교하면 끼워 넣은 줄 아래가 전부 '수정' 으로 나온다.

    8000줄짜리에서 줄 하나 넣고 저장하면 '7999줄 수정' 이라고 뜨는 셈이라,
    사람이 그 창을 안 읽게 된다 -- 읽으라고 띄우는 창인데.
    """
    before = frame([["1", "x"], ["2", "y"], ["3", "z"]])
    after = frame([["1", "x"], ["새", "줄"], ["2", "y"], ["3", "z"]])
    assert kinds(before, after) == ([("신규", 2)], 1)


def test_a_deleted_row_is_called_deleted():
    before = frame([["1", "x"], ["2", "y"], ["3", "z"]])
    after = frame([["1", "x"], ["3", "z"]])
    assert kinds(before, after) == ([("삭제", 2)], 1)


def test_nothing_changed_means_nothing_to_show():
    before = frame([["1", "x"], ["2", "y"]])
    assert kinds(before, before.copy()) == ([], 0)


def test_a_blank_row_left_over_is_not_a_change():
    """'행 아래' 를 눌렀다 안 채우고 저장하는 일이 흔하다."""
    before = frame([["1", "x"]])
    after = frame([["1", "x"], ["", ""]])
    assert kinds(before, after) == ([], 0)


def test_the_changed_row_carries_its_whole_row():
    before = frame([["1", "x"]])
    after = frame([["1", "바뀜"]])
    rows, _ = im.row_changes(before, after)
    assert rows[0]["values"] == {"a": "1", "b": "바뀜"}


def test_too_many_changes_are_counted_but_not_all_listed():
    before = frame([[str(i), "x"] for i in range(500)])
    after = frame([[str(i), "y"] for i in range(500)])
    rows, total = im.row_changes(before, after, limit=10)
    assert total == 500 and len(rows) == 10


def test_workbook_changes_leaves_out_untouched_sheets():
    before = {"A": frame([["1", "x"]]), "B": frame([["1", "x"]])}
    after = {"A": frame([["1", "바뀜"]]), "B": frame([["1", "x"]])}
    got = im.workbook_changes(before, after)
    assert list(got) == ["A"]
    assert got["A"]["total"] == 1


def test_workbook_changes_notices_a_new_column():
    before = {"A": frame([["1", "x"]])}
    after = {"A": frame([["1", "x", ""]], cols=("a", "b", "c"))}
    assert "칸 추가: c" in im.workbook_changes(before, after)["A"]["note"]


def test_workbook_changes_notices_a_dropped_sheet():
    got = im.workbook_changes({"A": frame([["1", "x"]]), "B": frame([["1", "x"]])},
                              {"A": frame([["1", "x"]])})
    assert got["B"]["note"] == "시트 삭제"


# ------------------------------------------------------------ REV_INFO

def rev_book():
    return {"STEP": frame([["1", "x"]]),
            "REV_INFO": pd.DataFrame(
                [{"Date": "2026-09-01", "Remark": "최초", "user": "hong", "관련": ""}],
                dtype=object)}


def test_rev_columns_are_found():
    assert im.rev_columns(rev_book()) == ["Date", "Remark", "user", "관련"]


def test_no_rev_sheet_means_no_reason_is_demanded():
    assert im.rev_columns({"STEP": frame([["1", "x"]])}) is None


def test_the_reason_is_appended_at_the_bottom():
    """위에 끼워 넣으면 다음에 열었을 때 그 시트의 줄 번호가 전부 밀린다."""
    book = rev_book()
    out = im.append_rev_info(book, "2026-09-21", "오탈자", "김철수", "JIRA-1")
    assert out["REV_INFO"].values.tolist() == [
        ["2026-09-01", "최초", "hong", ""],
        ["2026-09-21", "오탈자", "김철수", "JIRA-1"],
    ]


def test_appending_does_not_touch_what_was_passed_in():
    book = rev_book()
    im.append_rev_info(book, "2026-09-21", "사유", "나", "")
    assert len(book["REV_INFO"]) == 1


def test_the_optional_field_may_be_empty():
    out = im.append_rev_info(rev_book(), "2026-09-21", "사유", "나", "")
    assert out["REV_INFO"].iloc[-1]["관련"] == ""


def test_columns_are_matched_ignoring_case_and_spaces():
    book = {"REV_INFO": pd.DataFrame(columns=[" DATE ", "remark", "USER", "관련"],
                                     dtype=object)}
    out = im.append_rev_info(book, "2026-09-21", "사유", "나", "세부")
    assert out["REV_INFO"].iloc[-1].tolist() == ["2026-09-21", "사유", "나", "세부"]


def test_a_rev_sheet_with_other_columns_leaves_them_blank():
    book = {"REV_INFO": pd.DataFrame(columns=["Date", "Remark", "user", "관련", "기타"],
                                     dtype=object)}
    out = im.append_rev_info(book, "2026-09-21", "사유", "나", "")
    assert out["REV_INFO"].iloc[-1]["기타"] == ""


def test_the_saved_file_carries_the_new_rev_row():
    book = rev_book()
    with_rev = im.append_rev_info(book, "2026-09-21", "오탈자", "김철수", "")
    im.save_workbook("A", with_rev, "김철수")
    back = im.load_workbook("A")[0]["REV_INFO"]
    assert back.iloc[-1]["Remark"] == "오탈자"
    assert back.iloc[-1]["Date"] == "2026-09-21"


# ------------------------------------------------------------ 수식 지키기

def book_with_formula():
    """C 칸에 VLOOKUP 이 걸린 시트."""
    sheets = {"STEP": pd.DataFrame(
        [{"step": "0010", "ppid": "P-01", "찾은값": "가"},
         {"step": "0020", "ppid": "P-02", "찾은값": "나"}], dtype=object)}
    formulas = {"STEP": {(0, "찾은값"): "VLOOKUP(B2,MAP!$A$1:$B$99,2,0)",
                         (1, "찾은값"): "VLOOKUP(B3,MAP!$A$1:$B$99,2,0)"}}
    return sheets, formulas


def test_a_formula_survives_a_round_trip():
    """이게 안 되면 VLOOKUP 이 마지막 계산값으로 굳어 버린다."""
    sheets, formulas = book_with_formula()
    im.save_workbook("A", sheets, "hong", formulas=formulas)
    got: dict = {}
    im.load_workbook("A", got)
    assert got["STEP"][(0, "찾은값")] == "VLOOKUP(B2,MAP!$A$1:$B$99,2,0)"


def test_the_saved_file_really_holds_a_formula_not_a_value():
    sheets, formulas = book_with_formula()
    im.save_workbook("A", sheets, "hong", formulas=formulas)
    raw = fake_s3.STORE["2GAPU/input/A.xlsx"]
    book = openpyxl.load_workbook(io.BytesIO(raw))
    assert book["STEP"]["C2"].value == "=VLOOKUP(B2,MAP!$A$1:$B$99,2,0)"


def test_editing_another_column_keeps_the_formula():
    sheets, formulas = book_with_formula()
    after = {"STEP": sheets["STEP"].copy()}
    after["STEP"].loc[0, "ppid"] = "P-99"
    kept, lost = im.surviving_formulas(sheets, after, formulas)
    assert len(kept["STEP"]) == 2 and not lost


def test_typing_over_a_formula_cell_wins():
    """사람이 그 칸에 직접 값을 적었으면 그 값이 이긴다."""
    sheets, formulas = book_with_formula()
    after = {"STEP": sheets["STEP"].copy()}
    after["STEP"].loc[0, "찾은값"] = "손으로 적음"
    kept, lost = im.surviving_formulas(sheets, after, formulas)
    assert (0, "찾은값") not in kept.get("STEP", {})
    assert (1, "찾은값") in kept["STEP"]
    assert lost == {"STEP": 1}


def test_a_row_inserted_above_drops_the_formulas_below():
    """5번 줄의 =VLOOKUP(B5,..) 는 6번 줄로 밀리면 B6 을 봐야 맞다.

    자리를 따라 고쳐 주지는 못하므로, 틀린 수식을 남기는 대신 버린다.
    """
    sheets, formulas = book_with_formula()
    grown = pd.concat([
        pd.DataFrame([{"step": "0005", "ppid": "P-00", "찾은값": ""}], dtype=object),
        sheets["STEP"]], ignore_index=True)
    kept, lost = im.surviving_formulas(sheets, {"STEP": grown}, formulas)
    assert not kept
    assert lost == {"STEP": 2}


def test_a_row_added_at_the_end_keeps_the_formulas_above():
    sheets, formulas = book_with_formula()
    grown = pd.concat([
        sheets["STEP"],
        pd.DataFrame([{"step": "0030", "ppid": "P-03", "찾은값": ""}], dtype=object)],
        ignore_index=True)
    kept, lost = im.surviving_formulas(sheets, {"STEP": grown}, formulas)
    assert len(kept["STEP"]) == 2 and not lost


def test_a_dropped_sheet_drops_its_formulas():
    sheets, formulas = book_with_formula()
    kept, lost = im.surviving_formulas(sheets, {}, formulas)
    assert not kept and lost == {"STEP": 2}


def test_a_column_inserted_to_the_left_does_not_move_the_formula():
    """칸을 번호가 아니라 이름으로 붙들어 두는 자리다."""
    sheets, formulas = book_with_formula()
    after = sheets["STEP"].copy()
    after.insert(0, "새칸", ["", ""])
    kept, _lost = im.surviving_formulas(sheets, {"STEP": after}, formulas)
    im.save_workbook("A", {"STEP": after}, "hong", formulas=kept)
    raw = fake_s3.STORE["2GAPU/input/A.xlsx"]
    book = openpyxl.load_workbook(io.BytesIO(raw))
    assert book["STEP"]["A1"].value == "새칸"
    assert str(book["STEP"]["D2"].value).startswith("=VLOOKUP")


def test_a_blank_row_removed_on_save_pulls_the_formula_up():
    """_clean 이 빈 줄을 빼면 아래 줄이 당겨진다. 수식도 같이 당겨야 한다."""
    sheets = {"STEP": pd.DataFrame(
        [{"a": "", "f": ""}, {"a": "x", "f": "1"}], dtype=object)}
    formulas = {"STEP": {(1, "f"): "SUM(A3:A3)"}}
    im.save_workbook("A", sheets, "hong", formulas=formulas)
    raw = fake_s3.STORE["2GAPU/input/A.xlsx"]
    book = openpyxl.load_workbook(io.BytesIO(raw))
    assert book["STEP"]["B2"].value == "=SUM(A3:A3)", \
        [c.value for c in book["STEP"]["B"]]


def test_excel_is_told_to_recompute_on_open():
    """적어 둔 값은 낡았을 수 있다. 엑셀이 열 때 다시 계산해야 한다."""
    sheets, formulas = book_with_formula()
    im.save_workbook("A", sheets, "hong", formulas=formulas)
    raw = fake_s3.STORE["2GAPU/input/A.xlsx"]
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        assert b'fullCalcOnLoad="1"' in zf.read("xl/workbook.xml")


# ------------------------------------------------------------- 이력 내용



def test_the_change_text_reads_like_the_popup():
    """창에서 본 것과 같은 것이 파일에 남아야 한다 -- 나중에 되짚는 사람은
    그 창을 못 보고 이 칸만 본다."""
    before = {"STEP": frame([["1", "x"], ["2", "y"]])}
    after = {"STEP": frame([["1", "바뀜"], ["2", "y"], ["3", "새"]])}
    text = im.changes_text(im.workbook_changes(before, after))
    assert "[STEP]" in text
    assert "수정 1행: 1 | 바뀜" in text, text
    assert "신규 3행: 3 | 새" in text, text


def test_the_change_text_names_added_columns():
    before = {"STEP": frame([["1", "x"]])}
    after = {"STEP": frame([["1", "x", ""]], cols=("a", "b", "c"))}
    assert "칸 추가: c" in im.changes_text(im.workbook_changes(before, after))


def test_the_change_text_cannot_overflow_an_excel_cell():
    """엑셀 칸 하나는 32,767자까지다. 넘기면 파일이 안 열린다.

    줄 수는 300개로 막아 두지만 칸이 넓고 값이 길면 그것만으로 넘길 수 있다.
    """
    wide = [[("가" * 200) for _ in range(20)] for _ in range(300)]
    before = {"STEP": frame(wide, cols=[f"c{i}" for i in range(20)])}
    after = {"STEP": frame([[v + "!" for v in row] for row in wide],
                           cols=[f"c{i}" for i in range(20)])}
    text = im.changes_text(im.workbook_changes(before, after))
    assert len(text) <= im.CELL_MAX
    assert "잘림" in text, "잘렸다는 말이 없으면 뒤가 없는 건지 알 수 없다"


def test_a_truncated_change_text_says_so():
    before = {"STEP": frame([["1", "x"], ["2", "y"]])}
    after = {"STEP": frame([["1", "바뀜"], ["2", "또바뀜"]])}
    text = im.changes_text(im.workbook_changes(before, after), limit=40)
    assert len(text) <= 40 and text.endswith("잘림)"), repr(text)


def test_what_the_person_typed_comes_before_the_auto_part():
    out = im.append_rev_info(rev_book(), "2026-09-21", "사유", "나",
                             "JIRA-1", "[STEP] ...\n수정 1행: x")
    related = out["REV_INFO"].iloc[-1]["관련"]
    assert related.startswith("JIRA-1"), related
    assert "수정 1행" in related


def test_the_auto_part_alone_is_fine_when_nothing_was_typed():
    out = im.append_rev_info(rev_book(), "2026-09-21", "사유", "나", "",
                             "[STEP] ...")
    assert out["REV_INFO"].iloc[-1]["관련"] == "[STEP] ..."


# --------------------------------------------- 이력 폴더에 사본 쌓기

def hist(book="A"):
    return sorted(k.rsplit("/", 1)[-1]
                  for k in fake_s3.STORE if "/이력/" in k)


def test_every_save_drops_a_copy_in_the_history_folder():
    im.save_workbook("FAB_INPUT_ULY_r0", sheets(S=[{"a": 1}]), "junwoo.hwang")
    got = hist()
    assert len(got) == 1
    assert "_FAB_INPUT_ULY_r0_" in got[0], got
    assert got[0].endswith("_junwoo.hwang.xlsx"), got


def test_the_folder_sorts_by_time_because_the_date_comes_first():
    """이름순으로 봐도 시간순이 되게 날짜를 앞에 둔다."""
    import datetime as dt
    today = dt.datetime.now(im.KST).strftime("%y%m%d")
    im.save_workbook("Z파일", sheets(S=[{"a": 1}]), "hong")
    im.save_workbook("A파일", sheets(S=[{"a": 1}]), "kim")
    assert all(k.startswith(today) for k in hist()), hist()


def test_the_copy_is_named_by_the_day_and_the_person():
    import datetime as dt
    today = dt.datetime.now(im.KST).strftime("%y%m%d")
    im.save_workbook("A", sheets(S=[{"a": 1}]), "junwoo.hwang")
    assert hist() == [f"{today}_A_junwoo.hwang.xlsx"], hist()


def test_the_copy_holds_what_was_just_saved():
    im.save_workbook("A", sheets(S=[{"a": "새값"}]), "hong")
    key = next(k for k in fake_s3.STORE if "/이력/" in k)
    assert im.read_xlsx(fake_s3.STORE[key])["S"]["a"].tolist() == ["새값"]


def test_two_saves_on_the_same_day_by_the_same_person_do_not_collide():
    """그대로 두면 앞의 판이 조용히 덮어써진다 -- 되돌릴 것이 하나 사라진다."""
    im.save_workbook("A", sheets(S=[{"a": "첫째"}]), "hong")
    im.save_workbook("A", sheets(S=[{"a": "둘째"}]), "hong")
    im.save_workbook("A", sheets(S=[{"a": "셋째"}]), "hong")
    got = hist()
    assert len(got) == 3, got
    kept = sorted(im.read_xlsx(fake_s3.STORE[k])["S"]["a"][0]
                  for k in fake_s3.STORE if "/이력/" in k)
    assert kept == ["둘째", "셋째", "첫째"], kept


def test_the_history_copy_does_not_show_up_as_a_workbook():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "hong")
    assert im.list_workbooks() == ["A"]


def test_a_weird_user_id_cannot_escape_the_history_folder():
    im.save_workbook("A", sheets(S=[{"a": 1}]), "../../etc/passwd")
    keys = [k for k in fake_s3.STORE if "/이력/" in k]
    assert len(keys) == 1
    assert "/이력/" in keys[0] and "_A_" in keys[0], keys
    assert ".." not in keys[0], keys


def test_no_copy_no_overwrite():
    """사본을 못 남기면 본 파일도 안 건드린다."""
    im.save_workbook("A", sheets(S=[{"a": "원래"}]), "hong")
    before = fake_s3.STORE["2GAPU/input/A.xlsx"]
    real = fake_s3.put_object

    def no_history(key, data):
        if "/이력/" in key:
            raise RuntimeError("이력 폴더에 못 씁니다")
        return real(key, data)

    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(fake_s3, "put_object", no_history)
        with pytest.raises(RuntimeError):
            im.save_workbook("A", sheets(S=[{"a": "새것"}]), "hong")
    assert fake_s3.STORE["2GAPU/input/A.xlsx"] == before
    assert im.load_workbook("A")[0]["S"]["a"].tolist() == ["원래"]


def test_downloading_does_not_touch_the_history_folder():
    im.to_xlsx(sheets(S=[{"a": 1}]))
    assert not [k for k in fake_s3.STORE if "/이력/" in k]


# --------------------------------------------- VLOOKUP 을 대신 계산한다

def two_sheets():
    """B 가 ET추출여부 를 정확매칭 VLOOKUP 으로 끌어다 쓴다."""
    et = pd.DataFrame([
        {"코드": "P-01", "b": "x", "이름": "가"},
        {"코드": "P-02", "b": "y", "이름": "나"},
    ], dtype=object)
    b = pd.DataFrame({f"c{i}": [""] * 2 for i in range(5)}, dtype=object)
    b.loc[0, "c4"] = "P-01"
    b.loc[1, "c4"] = "P-02"
    formulas = {"B": {(0, "c0"): "VLOOKUP(E2,ET추출여부!$A:$C,3,0)",
                      (1, "c0"): "VLOOKUP(E3,ET추출여부!$A:$C,3,0)"}}
    return {"ET추출여부": et, "B": b}, formulas


def test_a_lookup_sheet_edit_updates_the_other_sheets_cached_value():
    """A 를 고쳐 저장하면, pandas 로 그 파일을 직접 읽는 쪽도 새 값을
    받아야 한다 -- 엑셀로 열어야만 맞는 값을 준다면 파이프라인이 못 쓴다."""
    sheets, formulas = two_sheets()
    sheets["ET추출여부"].loc[0, "이름"] = "새이름"
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == "새이름"
    assert got["B"].loc[1, "c0"] == "나"          # 안 건드린 줄은 그대로


def test_no_match_gives_the_excel_error_not_a_stale_value():
    sheets, formulas = two_sheets()
    sheets["B"].loc[0, "c4"] = "없는코드"
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == "#N/A"


def test_approximate_match_is_left_alone():
    """근사매칭(정렬을 가정하는 이분 탐색)은 계산을 흉내 내다 잘못 계산할
    위험이 크다 -- 안 하느니만 못하다. 손대지 않는다."""
    sheets, formulas = two_sheets()
    formulas["B"] = {(0, "c0"): "VLOOKUP(E2,ET추출여부!$A:$C,3,TRUE)"}
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == ""      # '가' 로 계산돼 있으면 안 된다


def test_other_functions_are_left_alone():
    sheets, formulas = two_sheets()
    formulas["B"] = {(0, "c0"): "SUMIF(A:A,\"x\",B:B)"}
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == sheets["B"].loc[0, "c0"] == ""


def test_the_recalculated_value_is_what_gets_saved():
    """실제 저장 경로로: A 를 고쳐 저장하면 파일에 남는 캐시가 새 값이다."""
    sheets, formulas = two_sheets()
    sheets["ET추출여부"].loc[0, "이름"] = "새이름"
    im.save_workbook("T", sheets, "hong", formulas=formulas)
    raw = fake_s3.STORE["2GAPU/input/T.xlsx"]
    cached = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)["B"]["A2"].value
    formula = openpyxl.load_workbook(io.BytesIO(raw))["B"]["A2"].value
    assert cached == "새이름"
    assert formula == "=VLOOKUP(E2,ET추출여부!$A:$C,3,0)"


def test_the_portal_itself_shows_the_fresh_value_after_saving():
    sheets, formulas = two_sheets()
    sheets["ET추출여부"].loc[0, "이름"] = "새이름"
    im.save_workbook("T", sheets, "hong", formulas=formulas)
    back = im.load_workbook("T")[0]
    assert back["B"]["c0"].tolist() == ["새이름", "나"]


def test_a_whole_column_range_and_a_bounded_range_both_work():
    sheets, formulas = two_sheets()
    formulas["B"][(0, "c0")] = "VLOOKUP(E2,ET추출여부!$A$1:$C$99,3,0)"
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == "가"


def test_looking_up_in_a_sheet_that_does_not_exist_is_left_alone():
    sheets, formulas = two_sheets()
    formulas["B"] = {(0, "c0"): "VLOOKUP(E2,없는시트!$A:$C,3,0)"}
    got = im.refresh_formula_cache(sheets, formulas)
    assert got["B"].loc[0, "c0"] == sheets["B"].loc[0, "c0"] == ""


def test_the_users_actual_formula_shapes_all_recompute():
    """실제로 받은 네 수식 -- 같은 시트, 같은 범위, 자기 줄의 E열을 키로
    쓰고 열번호(2 또는 3)와 적힌 자리(A 또는 C)만 다르다."""
    et = pd.DataFrame([
        {"코드": "P-01", "b값": "b2", "이름": "이름2"},
        {"코드": "P-02", "b값": "b3", "이름": "이름3"},
    ], dtype=object)
    main = pd.DataFrame({"A": ["", ""], "B": ["x", "y"], "C": ["", ""],
                         "D": ["p", "q"], "E": ["P-01", "P-02"]}, dtype=object)
    formulas = {"Main": {
        (0, "C"): "VLOOKUP(E2,ET추출여부!$A:$C,3,0)",
        (0, "A"): "VLOOKUP(E2,ET추출여부!$A:$C,2,0)",
        (1, "A"): "VLOOKUP(E3,ET추출여부!$A:$C,2,0)",
        (1, "C"): "VLOOKUP(E3,ET추출여부!$A:$C,3,0)",
    }}
    et2 = et.copy()
    et2.loc[0, "이름"] = "새이름2"
    et2.loc[1, "b값"] = "새b3"
    got = im.refresh_formula_cache({"ET추출여부": et2, "Main": main}, formulas)
    assert got["Main"].loc[0, "C"] == "새이름2"
    assert got["Main"].loc[0, "A"] == "b2"          # 안 바뀐 값은 그대로
    assert got["Main"].loc[1, "A"] == "새b3"
    assert got["Main"].loc[1, "C"] == "이름3"


def test_a_blank_gap_above_the_formula_does_not_confuse_the_cached_value():
    """빈 줄이 위에서 빠지면 그 아래 줄들이 저장 파일에서 자리가 당겨진다
    (_clean 이 늘 그래 왔다). 캐시 값은 그래도 맞아야 한다 -- 자리가
    당겨지기 전의 원래 자리를 기준으로 계산하기 때문이다.
    """
    # DataFrame 의 칸은 실제 엑셀 열(A,B,C,D,E...) 과 자리가 그대로
    # 맞아야 한다 -- 읽을 때 건너뛴 열은 'Unnamed: N' 으로 자리를 채워 두는
    # 것이 그래서다.
    et = pd.DataFrame([{"코드": "P-01", "b값": "b", "이름": "이름1"}], dtype=object)
    main = pd.DataFrame({
        "A": ["", "", ""], "B": ["", "", "y"], "C": ["", "", ""], "D": ["", "", ""],
        "E": ["", "", "P-01"],                   # 첫 두 줄은 통째로 빈 줄
    }, dtype=object)
    formulas = {"Main": {(2, "A"): "VLOOKUP(E4,ET추출여부!$A:$C,3,0)"}}
    got = im.refresh_formula_cache({"ET추출여부": et, "Main": main}, formulas)
    assert got["Main"].loc[2, "A"] == "이름1"


def test_the_first_matching_row_wins_like_excel():
    """같은 키가 두 줄이면 엑셀은 맨 위엣것을 준다."""
    et = pd.DataFrame([
        {"코드": "P-01", "b값": "b1", "이름": "위"},
        {"코드": "P-01", "b값": "b2", "이름": "아래"},
    ], dtype=object)
    main = pd.DataFrame({"A": [""], "B": [""], "C": [""], "D": [""],
                         "E": ["P-01"]}, dtype=object)
    formulas = {"Main": {(0, "C"): "VLOOKUP(E2,ET추출여부!$A:$C,3,0)"}}
    got = im.refresh_formula_cache({"ET추출여부": et, "Main": main}, formulas)
    assert got["Main"].loc[0, "C"] == "위"


def test_a_sheet_fixed_first_is_looked_up_with_its_new_values():
    """A 시트의 수식을 먼저 계산해 값이 갈렸으면, 그 시트를 찾아보는 B 시트는
    갈린 값을 봐야 한다. 찾을 표를 뒤집어 두고 쓰기 때문에 눌러 두는 자리다.
    """
    a = pd.DataFrame({"A": ["K1"], "B": ["헌값"], "C": [""]}, dtype=object)
    b = pd.DataFrame({"A": ["K1"], "B": ["새값"], "C": [""]}, dtype=object)
    main = pd.DataFrame({"A": [""], "B": [""], "C": [""], "D": [""],
                         "E": ["K1"]}, dtype=object)
    formulas = {
        # A 시트가 제 B 칸을 다른 시트에서 끌어와 채운다
        "A시트": {(0, "B"): "VLOOKUP(A2,원본!$A:$B,2,0)"},
        # Main 은 그 A 시트를 찾아본다
        "Main": {(0, "C"): "VLOOKUP(E2,A시트!$A:$B,2,0)"},
    }
    got = im.refresh_formula_cache(
        {"원본": b, "A시트": a, "Main": main}, formulas)
    assert got["A시트"].loc[0, "B"] == "새값"
    assert got["Main"].loc[0, "C"] == "새값", "낡은 값을 그대로 들고 왔습니다"


def test_recalculating_a_whole_sheet_of_formulas_is_not_quadratic():
    """줄마다 VLOOKUP 이 걸린 시트가 진짜로 있다. 수식마다 찾을 표를 처음부터
    훑으면 저장이 분 단위로 걸린다 -- 그래서 표를 한 번만 뒤집어 둔다.
    """
    import time
    look = pd.DataFrame({"A": [f"K{i}" for i in range(3000)],
                         "B": [f"b{i}" for i in range(3000)]}, dtype=object)
    n = 3000
    main = pd.DataFrame({"A": [""] * n, "B": [""] * n, "C": [""] * n,
                         "D": [""] * n, "E": [f"K{i}" for i in range(n)]},
                        dtype=object)
    formulas = {"Main": {(r, "A"): f"VLOOKUP(E{r + 2},찾을표!$A:$B,2,0)"
                         for r in range(n)}}
    start = time.perf_counter()
    got = im.refresh_formula_cache({"찾을표": look, "Main": main}, formulas)
    spent = time.perf_counter() - start
    assert got["Main"].loc[n - 1, "A"] == f"b{n - 1}"
    # 표를 한 번만 뒤집으면 0.2초쯤, 수식마다 훑으면 3초가 넘는다.
    assert spent < 1.5, f"{spent:.1f}초 걸렸습니다 -- 찾을 표를 또 훑고 있습니다"


# ------------------------------------------- 저장 뒤 raw data 반영 코드 돌리기

@pytest.fixture
def after_script(tmp_path, monkeypatch):
    """받은 인자를 파일에 적고 잠깐 도는 가짜 after_save.py."""
    out = tmp_path / "ran.txt"
    script = tmp_path / "after_save.py"
    script.write_text(
        "import sys, time\n"
        f"open({str(out)!r}, 'a', encoding='utf-8').write(' '.join(sys.argv[1:]) + '\\n')\n"
        "print('반영 중', sys.argv[1])\n"
        "time.sleep(float(__import__('os').environ.get('AFTER_SLEEP', '0')))\n",
        encoding="utf-8")
    monkeypatch.setattr(im, "AFTER_SAVE_SCRIPT", script)
    monkeypatch.setattr(im, "AFTER_SAVE_LOG", tmp_path / "after.log")
    im._after.clear()
    yield out
    im._after.clear()


def _wait_done(book, limit=20.0):
    import time
    end = time.time() + limit
    while time.time() < end:
        with im._after_lock:
            slot = im._after.get(book)
            if slot is None or slot["proc"] is None:
                return
        time.sleep(0.05)
    raise AssertionError("after_save 가 안 끝났습니다")


def test_the_raw_data_job_gets_the_file_the_path_the_user_and_the_version(after_script):
    im.run_after_save("A", "hong", "etag1")
    _wait_done("A")
    assert after_script.read_text(encoding="utf-8").split() == [
        "A", "2GAPU/input/A.xlsx", "hong", "etag1"]


def test_the_screen_does_not_wait_for_the_raw_data_job(after_script, monkeypatch):
    """20분 걸리는 코드다. 저장 완료는 바로 떠야 한다."""
    import time
    monkeypatch.setenv("AFTER_SLEEP", "3")
    start = time.time()
    im.run_after_save("A", "hong", "etag1")
    assert time.time() - start < 1.0
    _wait_done("A")


def test_saves_during_a_run_fold_into_one_more_run_with_the_latest(after_script,
                                                                    monkeypatch):
    """겹쳐 돌면 같은 raw data 를 둘이 동시에 쓴다. 도는 중에 몇 번을
    저장했든 끝난 뒤 한 번만, 가장 최근 판으로 더 돈다."""
    monkeypatch.setenv("AFTER_SLEEP", "1.5")
    im.run_after_save("A", "hong", "v1")
    im.run_after_save("A", "kim", "v2")
    im.run_after_save("A", "lee", "v3")
    _wait_done("A", limit=30)
    runs = after_script.read_text(encoding="utf-8").splitlines()
    assert [r.split()[3] for r in runs] == ["v1", "v3"], runs


def test_different_files_run_side_by_side(after_script, monkeypatch):
    monkeypatch.setenv("AFTER_SLEEP", "1")
    im.run_after_save("A", "hong", "a1")
    im.run_after_save("B", "hong", "b1")
    with im._after_lock:
        assert im._after["A"]["proc"] is not None
        assert im._after["B"]["proc"] is not None, "다른 파일까지 줄 세웠습니다"
    _wait_done("A"); _wait_done("B")


def test_what_the_job_prints_lands_in_the_log(after_script, tmp_path):
    im.run_after_save("A", "hong", "etag1")
    _wait_done("A")
    log = (tmp_path / "after.log").read_text(encoding="utf-8")
    assert "A 2GAPU/input/A.xlsx hong etag1" in log   # 언제 무엇으로 돌았나
    assert "반영 중 A" in log                          # 그 코드가 print 한 것


def test_the_empty_after_save_file_runs_cleanly(tmp_path, monkeypatch):
    """넣어 둔 빈 파일 그대로도 오류 없이 돌아야 한다."""
    import subprocess, sys
    got = subprocess.run([sys.executable, str(Path(im.__file__).with_name("after_save.py")),
                          "A", "2GAPU/input/A.xlsx", "hong", "etag"],
                         capture_output=True, text=True, timeout=30)
    assert got.returncode == 0, got.stderr
